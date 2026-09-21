"""Read the synthetic browser downloads with independent openpyxl assertions."""
from pathlib import Path
from io import BytesIO
import json
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / 'public/excel-converter/examples'
OUT = ROOT / 'outputs/excel-converter/business'
DOWNLOADS = Path.home() / 'Downloads'
FIXED = ['C3', 'C4', 'C6', 'C7', 'C8', 'C10', 'C11', 'C13', 'C14', 'C15', 'C17', 'C18', 'C19']

def equal_cells(actual, expected, addresses):
    for address in addresses:
        assert actual[address].value == expected[address].value, f'{address}: result differs from independent golden'
    return len(addresses)

def main():
    golden = openpyxl.load_workbook(SAMPLES / '12_위임장_정답예시.xlsx', data_only=True)
    with zipfile.ZipFile(DOWNLOADS / '변환완료_문서.zip') as archive:
        names = [name for name in archive.namelist() if name.endswith('.xlsx')]
        assert len(names) == 3
        delegation = 0
        for index, name in enumerate(names):
            sheet = openpyxl.load_workbook(BytesIO(archive.read(name)), data_only=True).worksheets[0]
            delegation += equal_cells(sheet, golden.worksheets[index], FIXED)
            assert len(sheet.merged_cells.ranges) == len(golden.worksheets[index].merged_cells.ranges)
    candidates = sorted(DOWNLOADS.glob('21_타회사노임명세서_한줄양식*변환*.xlsx'), key=lambda p:p.stat().st_mtime)
    assert candidates, 'Browser labor-statement download is missing'
    labor = openpyxl.load_workbook(candidates[-1], data_only=True).worksheets[0]
    expected = openpyxl.load_workbook(SAMPLES / '22_노임명세서_정답예시.xlsx', data_only=True).worksheets[0]
    addresses = ['B4'] + [f'{openpyxl.utils.get_column_letter(col)}{row}' for row in range(8, 12) for col in range(1, 13)]
    labor_count = equal_cells(labor, expected, addresses)
    assert labor['E11'].value == 7 and labor['G11'].value == 1270000
    formula_book = openpyxl.load_workbook(candidates[-1], data_only=False)
    assert formula_book.worksheets[0]['G8'].value == '=E8*F8'
    assert formula_book.worksheets[0]['G11'].value == '=SUM(G8:G10)'
    prompt_path = DOWNLOADS / '변환완료_문서 (1).zip'
    prompt_count = 0
    if prompt_path.exists():
        with zipfile.ZipFile(prompt_path) as archive:
            names = [name for name in archive.namelist() if name.endswith('.xlsx')]
            assert len(names) == 3
            for index, name in enumerate(names):
                sheet = openpyxl.load_workbook(BytesIO(archive.read(name)), data_only=True).worksheets[0]
                assert sheet['C4'].value == '2026-10'
                prompt_count += equal_cells(sheet, golden.worksheets[index], [a for a in FIXED if a != 'C4']) + 1
    report = {'browserDownloadChecks': {'delegationCells': delegation, 'laborCells': labor_count, 'promptCells': prompt_count}, 'mismatches': 0, 'people': 3, 'manDays': 7, 'amount': 1270000, 'preserved': ['merged cells', 'account leading zeros', 'labor formulas'], 'method': 'openpyxl compared browser downloads to independently authored golden workbooks'}
    (OUT / '브라우저_검증결과.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))

if __name__ == '__main__':
    main()
